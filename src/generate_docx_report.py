import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from pathlib import Path
import re

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>
            <w:insideV w:val="none"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

doc = docx.Document()

# Page Margins: 1 inch all around (standard ACM / academic)
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Set base styles
normal_style = doc.styles['Normal']
normal_style.font.name = 'Times New Roman'
normal_style.font.size = Pt(11)
normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
normal_style.paragraph_format.line_spacing = 1.15
normal_style.paragraph_format.space_after = Pt(4)

# 1. Header Metadata / Judul
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_after = Pt(8)
run_title = p_title.add_run("Keresahan Ketenagakerjaan dan Aspirasi Bekerja ke Luar Negeri di Era #KaburAjaDulu: Analisis Topik dan Sentimen Wacana Publik di Platform X")
run_title.font.name = 'Times New Roman'
run_title.font.size = Pt(16)
run_title.font.bold = True
run_title.font.color.rgb = RGBColor(0x11, 0x11, 0x11)

# Course subtitle
p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sub.paragraph_format.space_after = Pt(12)
run_sub = p_sub.add_run("Laporan Tugas Kelompok 2: Data Collection & Preparation Pipeline\nCSIE604284 Analitika Media Sosial — Semester Gasal 2026/2027")
run_sub.font.name = 'Times New Roman'
run_sub.font.size = Pt(11)
run_sub.font.italic = True
run_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

# Authors Table (2x2 grid for elegant layout)
table_auth = doc.add_table(rows=2, cols=2)
table_auth.alignment = WD_TABLE_ALIGNMENT.CENTER
table_auth.autofit = False

authors_data = [
    ("REGINA MEILANI ARUAN", "2306275632", "regina.meilani@ui.ac.id"),
    ("JOSHUA HANS VITHO SOEHENDRA", "2306165540", "joshua.hans@ui.ac.id"),
    ("ANGGUN SASMITHA DEWI", "2306165673", "anggun.sasmitha@ui.ac.id"),
    ("SALOMO IMMANUEL PUTRA", "2306219745", "salomo.immanuel@ui.ac.id")
]

idx = 0
for r in range(2):
    for c in range(2):
        name, npm, email = authors_data[idx]
        cell = table_auth.rows[r].cells[c]
        cell.width = Inches(3.2)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        r_name = p.add_run(f"{name}\n")
        r_name.font.bold = True
        r_name.font.size = Pt(10)
        r_det = p.add_run(f"NPM: {npm} | {email}\nProgram Studi Sistem Informasi\nFakultas Ilmu Komputer, Universitas Indonesia")
        r_det.font.size = Pt(9)
        r_det.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
        idx += 1

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# Abstrak Box
tbl_abs = doc.add_table(rows=1, cols=1)
tbl_abs.alignment = WD_TABLE_ALIGNMENT.CENTER
c_abs = tbl_abs.rows[0].cells[0]
c_abs.width = Inches(6.5)
set_cell_background(c_abs, "F8F9FA")
set_cell_margins(c_abs, top=140, bottom=140, left=200, right=200)

p_abs_h = c_abs.paragraphs[0]
p_abs_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_abs_h.paragraph_format.space_after = Pt(4)
r_abs_h = p_abs_h.add_run("ABSTRAK")
r_abs_h.font.bold = True
r_abs_h.font.size = Pt(11)

p_abs_b = c_abs.add_paragraph()
p_abs_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p_abs_b.paragraph_format.line_spacing = 1.15
p_abs_b.paragraph_format.space_after = Pt(6)
r_abs_b = p_abs_b.add_run(
    "Kondisi ketenagakerjaan di Indonesia menghadapi tantangan struktural yang signifikan, "
    "terutama tingginya angka pengangguran usia muda, ketidaksesuaian kualifikasi pendidikan "
    "dengan kebutuhan industri (overeducation), penyerapan sektor formal yang terbatas, dan "
    "gelombang pemutusan hubungan kerja (PHK). Dinamika ini memicu keresahan publik yang "
    "terartikulasi secara masif di media sosial, salah satunya melalui viralitas tagar #KaburAjaDulu "
    "di platform X yang mengaitkan krisis lapangan kerja domestik dengan aspirasi migrasi tenaga kerja "
    "ke luar negeri. Penelitian ini bertujuan untuk memetakan struktur tematik keresahan ketenagakerjaan "
    "dan menganalisis keterkaitannya dengan diskursus bekerja ke luar negeri beserta polaritas sentimen "
    "publik. Laporan Tugas Kelompok 2 (TK-2) ini mendokumentasikan implementasi end-to-end data pipeline, "
    "mulai dari pengumpulan digital trace data berbasis 36 kata kunci discovery, deduplikasi temporal "
    "WIB, penyaringan relevansi semantik hibrida (rule-hybrid dan audit manusia), pra-pemrosesan teks "
    "konservatif multi-representasi, hingga pembekuan korpus riset rilis final (tk2_v1) sebanyak 1.491 cuitan "
    "bersih. Seluruh pipeline terverifikasi lulus 46 pengujian otomatis (100% PASS), menjamin integritas "
    "dan keterulangan (reproducibility) data untuk tahapan topic modeling dan sentiment analysis berikutnya."
)
r_abs_b.font.size = Pt(9.5)

p_abs_kw = c_abs.add_paragraph()
p_abs_kw.paragraph_format.space_after = Pt(0)
r_kw_h = p_abs_kw.add_run("Kata Kunci: ")
r_kw_h.font.bold = True
r_kw_h.font.size = Pt(9.5)
r_kw_b = p_abs_kw.add_run("Ketenagakerjaan, #KaburAjaDulu, Platform X, Data Collection Pipeline, Text Preprocessing, Analitika Media Sosial.")
r_kw_b.font.size = Pt(9.5)
r_kw_b.font.italic = True

doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_heading_1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    return p

def add_heading_2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    return p

def add_body_p(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10.5)
    return p

# 1. PENDAHULUAN
add_heading_1("A. BAGIAN PENDAHULUAN")

# A1. Permasalahan Penelitian
add_heading_2("A1. Permasalahan Penelitian")
add_body_p(
    "Kondisi ketenagakerjaan di Indonesia menghadapi tekanan berat akibat tingginya tingkat "
    "pengangguran usia muda, menyusutnya penyerapan sektor formal, maraknya fenomena overeducation, "
    "dan berlanjutnya gelombang pemutusan hubungan kerja (PHK) [1], [2]. Tekanan ini menimbulkan "
    "kecemasan psikososial yang intens di kalangan pencari kerja [3]. Masalah utama yang diteliti adalah "
    "bagaimana keresahan ketenagakerjaan tersebut diekspresikan secara spontan oleh publik di platform X "
    "dan bagaimana wacana tersebut memicu aspirasi bekerja ke luar negeri melalui fenomena viral tagar "
    "#KaburAjaDulu [5], [6]."
)

# A2. Kesenjangan Penelitian
add_heading_2("A2. Kesenjangan Penelitian")
add_body_p(
    "Kajian ketenagakerjaan di Indonesia selama ini umumnya bertumpu pada survei statistik nasional "
    "yang tidak merekam keluhan spontan masyarakat [1], [2]. Sementara itu, studi media sosial di "
    "Indonesia masih terbatas pada klasifikasi sentimen biner terhadap PHK atau tagar tertentu tanpa "
    "memetakan struktur tematik keresahan secara mendalam [6], [10]. Penelitian ini mengisi kesenjangan "
    "tersebut dengan mengeksplorasi keterkaitan antara topik keresahan domestik dan aspirasi bekerja ke "
    "luar negeri menggunakan data media sosial X."
)

# A3. Tujuan Penelitian
add_heading_2("A3. Tujuan Penelitian")
add_body_p(
    "Penelitian ini bertujuan untuk memahami struktur tematik wacana publik di platform X mengenai "
    "keresahan ketenagakerjaan di Indonesia beserta keterkaitannya dengan aspirasi bekerja ke luar "
    "negeri dan polaritas sentimen publik pada periode fenomena #KaburAjaDulu."
)

# A4. Pertanyaan Penelitian
add_heading_2("A4. Pertanyaan Penelitian")
add_body_p(
    "Berdasarkan rumusan permasalahan dan kesenjangan penelitian di atas, pertanyaan penelitian dirumuskan sebagai berikut:\n"
    "1. Apa saja topik utama keresahan ketenagakerjaan yang dibicarakan dalam wacana publik di platform X?\n"
    "2. Bagaimana aspirasi bekerja ke luar negeri dibicarakan, dan bagaimana keterkaitannya dengan topik keresahan ketenagakerjaan domestik?\n"
    "3. Bagaimana distribusi sentimen publik pada setiap topik keresahan ketenagakerjaan dan aspirasi bekerja ke luar negeri?"
)

# A5. Pendekatan Analitika Media Sosial yang Digunakan
add_heading_2("A5. Pendekatan Analitika Media Sosial yang Digunakan")
add_body_p(
    "Penelitian ini menerapkan tiga teknik analitika media sosial. Pertama, descriptive analysis digunakan "
    "untuk memetakan karakteristik volume unggahan, pola temporal, metrik keterlibatan (likes, reposts), "
    "dan profil linguistik cuitan. Kedua, topic modeling berbasis BERTopic diterapkan untuk mengekstraksi "
    "kluster topik keresahan ketenagakerjaan dan diskursus bekerja ke luar negeri secara tanpa pengawasan "
    "(unsupervised). Ketiga, sentiment analysis digunakan untuk mengukur polaritas opini (positif, netral, "
    "negatif) pada masing-masing topik yang teridentifikasi. Kombinasi ini memungkinkan pengungkapan keterkaitan "
    "substantif antara dorongan krisis ketenagakerjaan lokal dengan aspirasi mobilitas internasional secara komprehensif."
)

# A6. Kontribusi Penelitian
add_heading_2("A6. Kontribusi Penelitian")
add_body_p(
    "Secara teoretis, penelitian ini memperluas literatur analitika media sosial di Indonesia dengan "
    "menggabungkan pemodelan topik dan analisis sentimen untuk menangkap dinamika pasar tenaga kerja. "
    "Secara praktis, temuan penelitian ini memberikan bukti empiris pelengkap (complementary digital evidence) "
    "bagi pembuat kebijakan ketenagakerjaan dalam memahami isu-isu paling mendesak yang dirasakan angkatan kerja "
    "muda, seperti upah minimum, kualifikasi lowongan, dan dorongan migrasi internasional."
)

# B. ULASAN LITERATUR
add_heading_1("B. ULASAN LITERATUR")
add_body_p(
    "Penelitian terdahulu menunjukkan bahwa persoalan ketenagakerjaan dapat dipahami melalui tiga perspektif utama. "
    "Pertama, kajian ekonometrika berbasis data Sakernas menegaskan bahwa tantangan pasar kerja Indonesia tidak "
    "hanya menyangkut kuantitas pengangguran terbuka, melainkan juga persoalan kualitas dan kesesuaian pekerjaan, "
    "terutama fenomena overeducation yang menimbulkan penalti upah yang signifikan bagi tenaga kerja terdidik [2]. "
    "Kedua, studi psikososial dan kualitatif mengenai transisi dari dunia pendidikan ke dunia kerja membuktikan "
    "bahwa para pencari kerja menghadapi kecemasan kerja yang mendalam serta hambatan struktural yang tidak "
    "sepenuhnya dapat ditangkap hanya melalui indikator ketersediaan lowongan pekerjaan semata [3], [8]. "
    "Ketiga, studi komputasional berbasis media sosial menunjukkan bahwa digital trace data sangat efektif dalam "
    "merefleksikan kondisi psikologis pengangguran [9], memetakan opini publik terhadap gelombang pemutusan hubungan kerja "
    "(PHK) massal [10], serta mengidentifikasi ekspresi kritik sosial kolektif terhadap kebijakan kesejahteraan publik "
    "pada tagar #KaburAjaDulu [6], [12]. Selain itu, literatur migrasi tenaga kerja internasional mengonfirmasi bahwa "
    "keresahan sosial-ekonomi domestik merupakan faktor pendorong (push factor) yang lebih dominan dalam memicu aspirasi "
    "bekerja ke luar negeri dibandingkan sekadar pengaruh paparan informasi media sosial [4], [5]. Kendati demikian, "
    "kajian media sosial ketenagakerjaan di Indonesia selama ini masih terfokus pada analisis sentimen agregat satu isu dan "
    "belum memetakan struktur tematik keresahan secara eksploratif. Penelitian ini mensintesis temuan-temuan tersebut dengan "
    "menganalisis percakapan publik di platform X menggunakan pendekatan terintegrasi antara pemodelan topik dan analisis sentimen."
)

# C. PENGUMPULAN DATA
add_heading_1("C. PENGUMPULAN DATA")

# C1. Kebutuhan Data
add_heading_2("C1. Kebutuhan Data")
add_body_p(
    "Untuk menjawab ketiga pertanyaan penelitian, kebutuhan data difokuskan pada digital trace data publik di platform X "
    "yang merekam secara spontan persepsi dan keresahan masyarakat mengenai kondisi ketenagakerjaan di Indonesia serta "
    "diskursus bekerja ke luar negeri. Platform X dipilih sebagai sumber data utama karena sifat percakapannya yang terbuka, "
    "berbasis teks ringkas, bertempo cepat, dan telah terbukti secara akademis sebagai wadah representatif bagi opini publik "
    "terkait isu perburuhan dan kebijakan ketenagakerjaan nasional [10]. Kebutuhan data mencakup unit analisis berupa teks unggahan "
    "utama (original tweets), balasan percakapan (replies), dan kutipan (quote posts) yang dipublikasikan dalam rentang waktu "
    "1 Januari 2026 hingga 25 Oktober 2026 pada zona waktu Asia/Jakarta (WIB). Metadata pendukung yang dikumpulkan meliputi "
    "stempel waktu pembuatan unggahan (createdAt), pengenal unik cuitan (id), metrik keterlibatan publik (likes, retweets, replies, quotes), "
    "serta atribut kebahasaan. Kriteria inklusi mensyaratkan teks berbahasa Indonesia atau campuran Indonesia-Inggris yang memuat kata "
    "kunci substantif ketenagakerjaan, upah, PHK, dan aspirasi bekerja ke luar negeri. Kriteria eksklusi menolak konten promosi komersial, "
    "iklan lowongan penipuan, penawaran jasa joki, berita olahraga, candaan tanpa konteks kerja, dan cuitan URL-only tanpa teks linguistik."
)

# Table 3.1: Pemetaan Kebutuhan Data terhadap Pertanyaan Penelitian
p_tbl_lbl = doc.add_paragraph()
p_tbl_lbl.paragraph_format.space_before = Pt(4)
p_tbl_lbl.paragraph_format.space_after = Pt(2)
r_tl = p_tbl_lbl.add_run("Tabel 3.1 Pemetaan Kebutuhan Data terhadap Pertanyaan Penelitian.")
r_tl.font.bold = True
r_tl.font.size = Pt(10)

tbl_rq = doc.add_table(rows=4, cols=4)
tbl_rq.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl_rq)

headers = ["Pertanyaan Penelitian", "Kebutuhan Data", "Atribut & Metadata Utama", "Justifikasi / Alasan Metodologis"]
widths = [Inches(1.2), Inches(1.6), Inches(1.7), Inches(2.0)]

for i, h in enumerate(headers):
    cell = tbl_rq.rows[0].cells[i]
    cell.width = widths[i]
    set_cell_background(cell, "EAECEF")
    set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(h)
    r.font.bold = True
    r.font.size = Pt(9.5)

rq_table_data = [
    (
        "RQ-1: Topik Keresahan Ketenagakerjaan",
        "Digital trace data cuitan publik mengenai pasar kerja, lowongan, pengangguran, upah, dan PHK di Indonesia.",
        "text_raw, text_topic, post_id, timestamp_wib, query_groups.",
        "Mengekstraksi representasi semantik leksikal untuk klasterisasi topik keresahan via BERTopic."
    ),
    (
        "RQ-2: Wacana & Aspirasi Luar Negeri",
        "Unggahan publik mengenai tagar #KaburAjaDulu, informasi kerja luar negeri, dan kaitan kausal domestik.",
        "text_raw, abroad_aspiration, abroad_discourse, issue_abroad_link.",
        "Membedakan aspirasi migrasi personal, wacana ketenagakerjaan global, dan relasi dorongan ekonomi lokal."
    ),
    (
        "RQ-3: Polaritas Sentimen Publik",
        "Cuitan dengan fitur sentimen terjaga (kata negasi, intensifier, tanda baca ekspresif, dan emotikon).",
        "text_raw, text_sentiment, like_count, retweet_count, reply_count.",
        "Mengukur valensi emosi dan tingkat polaritas publik terhadap masing-masing klaster topik persoalan."
    )
]

for row_idx, data in enumerate(rq_table_data, start=1):
    for col_idx, text in enumerate(data):
        cell = tbl_rq.rows[row_idx].cells[col_idx]
        cell.width = widths[col_idx]
        set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if col_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        r.font.size = Pt(9)

doc.add_paragraph().paragraph_format.space_after = Pt(4)

# C2. Pipeline Pengumpulan Data
add_heading_2("C2. Pipeline Pengumpulan Data")
add_body_p(
    "Pipeline pengumpulan dan persiapan data dirancang secara terstruktur dari ujung ke ujung (end-to-end), modular, "
    "dan dapat direproduksi sepenuhnya secara offline (reproducible) guna menjamin validitas metodologis riset. Desain pipeline "
    "mengintegrasikan enam tahapan utama yang saling berkesinambungan:\n\n"
    "1. Data Collection: Pengumpulan digital trace data dari platform X menggunakan crawler terkelola berbasis 36 kata kunci discovery, "
    "menghasilkan 3.292 rekaman JSONL mentah yang disimpan secara persisten dan immutable.\n"
    "2. Preparation & Deduplication: Penguraian stempel waktu ke zona waktu Asia/Jakarta (WIB), validasi kepatuhan jendela temporal mingguan, "
    "dan deduplikasi berbasis pengenal post_id unik untuk menghasilkan 1.781 kandidat unik.\n"
    "3. Semantic Relevance Screening: Penyaringan relevansi menggabungkan otoritas 95 label ground truth manual manusia (64 include, 31 exclude) "
    "dan model klasifikasi aturan semantik leksikal-diskursus lokal (anmedsos-rule-hybrid-v1) guna mengeliminasi spam, promosi, dan olahraga.\n"
    "4. Multi-Label RQ Classification: Pengklasifikasian teks ke dalam skema multi-label Research Questions (RQ1 Isu Kerja, RQ2 Aspirasi Personal, "
    "RQ2 Wacana Publik, dan RQ2 Kausal Eksplisit) dengan bukti cuplikan teks verbatim.\n"
    "5. Text Preprocessing: Normalisasi teks deterministik melalui pembersihan URL/mention, unescape entitas HTML, normalisasi Unicode NFC, "
    "casefolding, serta formalisasi bahasa gaul dan singkatan non-baku menggunakan kamus leksikal terstandarisasi.\n"
    "6. Quality Audit & Release Packaging: Evaluasi silang independen menggunakan model NLI transformer (mDeBERTa-v3-base XNLI), pemisahan satu cuitan "
    "URL-only ke exclusion log, dan pembekuan paket rilis korpus riset tk2_v1 sebanyak 1.491 baris bersih beserta manifest metadata lengkap."
)

# Textual Pipeline Diagram Box
p_diag_lbl = doc.add_paragraph()
p_diag_lbl.paragraph_format.space_before = Pt(4)
p_diag_lbl.paragraph_format.space_after = Pt(2)
r_dl = p_diag_lbl.add_run("Gambar 3.1 Diagram Alir End-to-End Data Collection & Preprocessing Pipeline ANMEDSOS.")
r_dl.font.bold = True
r_dl.font.size = Pt(10)

tbl_diag = doc.add_table(rows=1, cols=1)
tbl_diag.alignment = WD_TABLE_ALIGNMENT.CENTER
c_diag = tbl_diag.rows[0].cells[0]
c_diag.width = Inches(6.5)
set_cell_background(c_diag, "F4F6F8")
set_cell_margins(c_diag, top=100, bottom=100, left=140, right=140)

p_d = c_diag.paragraphs[0]
p_d.paragraph_format.line_spacing = 1.05
p_d.paragraph_format.space_after = Pt(0)
r_diag = p_d.add_run(
    "[1. RAW DATA COLLECTION]\n"
    "   Platform X via Apify Actor (36 Discovery Queries, Grup A–D)\n"
    "   Output: data/raw/raw_tweets.jsonl (3.292 JSON lines mentah, immutable)\n"
    "       │\n"
    "       ▼\n"
    "[2. PREPARATION & TEMPORAL DEDUPLICATION]\n"
    "   Konversi WIB, Validasi Interval Jendela Mingguan, Unik post_id\n"
    "   Output: data/interim/screened_candidates.csv (1.781 kandidat unik)\n"
    "       │\n"
    "       ▼\n"
    "[3. HYBRID SEMANTIC RELEVANCE SCREENING]\n"
    "   95 Audit Manual Manusia (Ground Truth) + Aturan Leksikal-Diskursus Lokal\n"
    "   Filter: Promosi Komersial, Olahraga, Berita Asing Non-Indonesia, Mention Insidental\n"
    "   Output: data/interim/semantic_screened.csv (1.516 model include + 1 recovered - 25 excluded)\n"
    "       │\n"
    "       ▼\n"
    "[4. MULTI-LABEL RQ CLASSIFICATION]\n"
    "   Pemetaan Morfologis & Leksikal ke RQ1, RQ2 Personal, RQ2 Publik, & RQ2 Kausal\n"
    "   Output: data/processed/final_corpus.csv (1.492 postingan korpus interim)\n"
    "       │\n"
    "       ▼\n"
    "[5. TEXT PREPROCESSING & QUALITY RECONCILIATION]\n"
    "   Unescape HTML, Unicode NFC, URL/Mention Stripping, Kamus Slang Formal (16 pasangan kata)\n"
    "   Evaluasi Truncation, Cross-Validation mDeBERTa-v3-base-xnli, Filter URL-only\n"
    "   Output: data/releases/tk2_v1/exclusion_log.csv (1 post URL-only disisihkan)\n"
    "       │\n"
    "       ▼\n"
    "[6. RESEARCH-READY DATASET RELEASE]\n"
    "   data/releases/tk2_v1/corpus.csv (1.491 postingan riset bersih)\n"
    "   data/releases/tk2_v1/preprocessed.csv (1.491 baris teks multi-representasi)\n"
    "   data/releases/tk2_v1/manifest.json (SHA-256 Hashes, Git Commit, & RQ Stats)"
)
r_diag.font.name = 'Courier New'
r_diag.font.size = Pt(8.5)

doc.add_paragraph().paragraph_format.space_after = Pt(4)

# C3. Metode Pengumpulan Data
add_heading_2("C3. Metode Pengumpulan Data")
add_body_p(
    "Pengumpulan data diimplementasikan secara teknis memanfaatkan crawler terkelola Apify (scrapesmith/x-twitter-search-scraper) "
    "yang beroperasi melalui protokol REST API dengan otentikasi Bearer Token terenkripsi pada environment lokal (.env). "
    "Strategi pengumpulan menerapkan kueri Boolean terencana yang mencakup 36 kata kunci pencarian discovery, terbagi ke dalam "
    "empat kelompok tematik: Grup A (pencarian kerja umum: 'cari kerja', 'lowongan kerja', 'fresh graduate'), Grup B (keluhan kerja dan "
    "pemutusan hubungan kerja: 'susah cari kerja', 'PHK', 'gaji rendah', 'pengangguran'), Grup C (wacana migrasi: '#KaburAjaDulu', "
    "'kerja di luar negeri', 'kerja di Jepang'), dan Grup D (istilah ketenagakerjaan bahasa Inggris: 'jobs in Indonesia', 'unemployment Indonesia'). "
    "Konstruksi kueri menggabungkan frasa pencarian leksikal dengan pembatas rentang waktu sejak (since) dan hingga (until) sesuai "
    "jadwal sampling mingguan serta mode pengurutan cuitan terkini (Latest). Pengambilan data mentah menerapkan penanganan pagination "
    "berbasis offset bertahap pada endpoint dataset Apify hingga seluruh item hasil pencarian berhasil diekstraksi ke penyimpanan lokal. "
    "Pemantauan eksekusi dilakukan melalui polling status run secara berkala dengan interval 5 detik hingga batas maksimum 120 iterasi. "
    "Setiap kegagalan tugas atau timeout dicatat secara transparan pada struktur log run (status ERROR) guna menjamin akuntabilitas audit "
    "tanpa mengorbankan integritas data historis. Seluruh rekaman mentah diekstraksi ke format JSON Lines pada data/raw/raw_tweets.jsonl yang "
    "memuat muatan lengkap, atribut penulis, metrik publik (likes, retweets, replies), dan stempel waktu UTC asli, menyediakan basis audit "
    "yang kuat bagi keseluruhan penelitian."
)

# C4. Implementasi Pipeline Pengumpulan Data & Preprocessing
add_heading_2("C4. Implementasi Pipeline Pengumpulan Data & Preprocessing")
add_body_p(
    "Implementasi pipeline pengumpulan dan persiapan data telah dieksekusi secara nyata dan berhasil menghasilkan dataset riset yang siap "
    "dianalisis. Dari 3.292 rekaman mentah yang terkumpul pada data/raw/raw_tweets.jsonl, modul prepare.py menyaring 1.781 kandidat unik "
    "yang valid secara temporal pada zona waktu WIB. Modul semantic_screen.py kemudian menerapkan penyaringan relevansi hibrida yang memprioritaskan "
    "95 keputusan audit manual manusia dan menyaring kandidat selebihnya menggunakan pengklasifikasi aturan leksikal-diskursus lokal, "
    "menghasilkan 1.492 postingan terpilih. Tahap pra-pemrosesan teks dieksekusi secara deterministik melalui modul preprocess_text.py yang "
    "menjalankan normalisasi Unicode NFC, unescape entitas HTML, penghapusan tautan URL dan mention pengguna (@username), casefolding, serta substitusi "
    "kata gaul dan singkatan non-baku menggunakan kamus formal 16 pasangan kata leksikal ketenagakerjaan (config/slang_dictionary.json). Tiga representasi teks dihasilkan "
    "secara konsisten untuk analisis bertahap: text_clean, text_topic, dan text_sentiment. Dalam audit kualitas pra-pemrosesan, ditemukan satu postingan URL-only "
    "(post_id: 2030131219137278032) yang menghasilkan string kosong setelah pembersihan tautan, sehingga dipisahkan secara transparan ke dalam "
    "data/releases/tk2_v1/exclusion_log.csv dengan alasan text_empty_after_preprocessing. Korpus akhir resmi dibekukan pada data/releases/tk2_v1/corpus.csv "
    "dan preprocessed.csv sebanyak 1.491 postingan riset bersih dengan cakupan multi-label yang solid (RQ1: 985 cuitan/66,06%, RQ2 Personal: 207 cuitan/13,88%, "
    "RQ2 Publik: 447 cuitan/29,98%, dan RQ2 Kausal: 8 cuitan/0,54%). Seluruh kode program, konfigurasi, dan test suite terorganisir rapi dan dapat "
    "diakses secara publik pada repositori GitHub resmi: https://github.com/rerearuan/kaburajadulu-analitik.git."
)

# C5. Kualitas Data dan Keterbatasannya
add_heading_2("C5. Kualitas Data dan Keterbatasannya")
add_body_p(
    "Potensi isu kualitas data telah diidentifikasi dan dimitigasi secara terstruktur. Isu duplikasi data diselesaikan "
    "melalui penegakan keunikan pengenal post_id (0 duplikasi). Isu missing data teratasi dengan mendeteksi dan memisahkan "
    "satu cuitan non-teks URL-only ke exclusion log, memastikan 100% korpus rilis memiliki teks bersih yang bermakna. "
    "Keterbatasan temporal tercatat secara transparan: meskipun desain awal menargetkan 1 Januari – 25 Oktober 2026, data aktual "
    "yang berhasil tersampel melalui jendela mingguan mencakup rentang 1 Januari – 7 Oktober 2026 (WIB), di mana periode 8–25 Oktober "
    "tidak tersampel oleh jadwal crawler. Evaluasi silang independen menggunakan model NLI transformer (mDeBERTa-v3-base-xnli) "
    "menghasilkan raw agreement 85,35%, namun mengonfirmasi adanya over-inclusion pada model deep learning ketika hipotesis komposit "
    "diuji pada cuitan di luar domain (misalnya berita regional Malaysia atau wacana kripto yang berhasil disaring oleh classifier awal). "
    "Korpus rilis berstatus provisional (64 postingan diaudit manual manusia dan 1.427 postingan terseleksi secara leksikal-diskursus)."
)

# C6. Pertimbangan Etika
add_heading_2("C6. Pertimbangan Etika")
add_body_p(
    "Penelitian ini mematuhi prinsip etika penelitian media sosial dan ketentuan X Developer Agreement & Policy. "
    "Pengumpulan data dibatasi secara ketat pada konten yang bersifat terbuka untuk umum (public posts); tidak ada data percakapan privat "
    "atau akun terkunci yang diakses. Perlindungan privasi pengguna diterapkan melalui penghapusan mention akun (@username) pada teks pra-pemrosesan "
    "serta pengecualian atribut pribadi sensitif (seperti nomor kontak atau alamat fisik). "
    "Guna mematuhi aturan platform mengenai larangan redistribusi publik payload teks mentah pengguna, repositori GitHub publik hanya "
    "mempublikasikan kode sumber, dokumentasi, file konfigurasi, laporan audit agregat, dan metadata rilis (manifest.json). "
    "Seluruh dataset teks mentah asli (raw_tweets.jsonl) dan korpus lengkap disimpan secara aman pada penyimpanan lokal terkontrol "
    "dan hanya diserahkan kepada dosen penilai melalui akses terbatas akademik."
)

# C7. Penutup
add_heading_2("C7. Penutup")
add_body_p(
    "Strategi pengumpulan data dan pra-pemrosesan teks yang dikembangkan dalam Tugas Kelompok 2 ini berhasil membangun "
    "fondasi empiris yang kokoh, bersih, dan berstandar ilmiah tinggi untuk mendukung pencapaian tujuan penelitian. "
    "Pipeline ini berhasil mentransformasikan 3.292 data mentah digital trace menjadi korpus riset terverifikasi sebanyak 1.491 postingan "
    "yang kaya konteks dan bebas noise. Pemetaan multi-label yang seimbang terhadap 985 cuitan isu ketenagakerjaan domestik (RQ1), "
    "207 cuitan aspirasi migrasi personal, dan 447 cuitan wacana publik luar negeri (RQ2) memastikan ketersediaan data yang sangat memadai "
    "untuk diekstraksi tema-temanya menggunakan BERTopic dan diukur valensi emosinya melalui sentiment analysis pada tahapan Tugas Kelompok berikutnya."
)

# REFERENCES
add_heading_1("REFERENCES")

references_list = [
    "[1] Tempo. 2025. BPS: 7.28 Million Indonesians Unemployed as of February 2025. Tempo. Retrieved from https://en.tempo.co/read/2003766/bps-7-28-million-indonesians-unemployed-as-of-february-2025",
    "[2] P. Wicaksono, I. Theresia, and B. Al Aufa. 2023. Education–Occupation Mismatch and Its Wage Penalties: Evidence from Indonesia. Cogent Business & Management 10, 3 (2023), 2251206. https://doi.org/10.1080/23311975.2023.2251206",
    "[3] E. Çarkıt. 2025. A Social Cognitive Lens on Employment Anxiety and Career Optimism in the Transition from School to Work. The Career Development Quarterly 73, 3 (2025), 194–204. https://doi.org/10.1002/cdq.70000",
    "[4] D. Adegoke. 2023. “Japa”: An Exploratory Study of the Roles of Social Media in an Out-Migration Trend in Nigeria. Social Media + Society 9, 4 (2023). https://doi.org/10.1177/20563051231203691",
    "[5] BenarNews. 2025. Indonesian youths look abroad for better economic future. BenarNews. Retrieved from https://www.benarnews.org/english/news/indonesian/youth-look-abroad-for-future-02192025140813.html",
    "[6] Z. Putri Ajie and A. Wardhana. 2025. Analisis Sentimen Negatif Publik pada Tagar #KaburAjaDulu di Media Sosial “X”. Jurnal Audiens 6, 3 (2025), 430–442. https://doi.org/10.18196/jas.v6i3.608",
    "[7] J. Ohme, T. Araujo, L. Boeschoten, D. Freelon, N. Ram, B. B. Reeves, and T. N. Robinson. 2024. Digital Trace Data Collection for Social Media Effects Research: APIs, Data Donation, and (Screen) Tracking. Communication Methods and Measures 18, 2 (2024), 124–141. https://doi.org/10.1080/19312458.2023.2181319",
    "[8] A. S. Wani, I. Joshi, and P. Singh. 2024. Navigating the Job-Seeking Journey: Challenges and Opportunities for Digital Employment Support in Kashmir. Proceedings of the ACM on Human-Computer Interaction 8, CSCW1 (2024), 1–28. https://doi.org/10.1145/3637375",
    "[9] M. Tan, Z. Wu, J. Li, Y. Liang, and W. Lv. 2024. Analyzing the Impact of Unemployment on Mental Health among Chinese University Graduates: A Study of Emotional and Linguistic Patterns on Weibo. Frontiers in Public Health 12 (2024), 1337859. https://doi.org/10.3389/fpubh.2024.1337859",
    "[10] A. M. AlHakim, P. D. Atika, and Herlawati. 2025. Analisis Sentimen Masyarakat Terhadap PHK di Indonesia Pada Twitter Menggunakan Naïve Bayes dan Support Vector Machine (SVM). Journal of Students' Research in Computer Science (JSRCS) 6, 1 (2025), 113–124. https://doi.org/10.31599/96sfw544",
    "[11] K. S. Rosamma. 2024. Analyzing Online Conversations on Reddit: A Study of Stress and Anxiety Through Topic Modeling and Sentiment Analysis. Cureus 16, 9 (2024), e69030. https://doi.org/10.7759/cureus.69030",
    "[12] T. Hamim Aftoni, W. H. Sugiharto, and M. I. Ghozali. 2025. Analisis Sentimen terhadap Tagar #KaburAjaDulu di Media Sosial X Menggunakan Metode Naïve Bayes. Jurnal Komputer dan Teknologi 4, 2 (2025), 154–167. https://doi.org/10.64626/jukomtek.v4i2.443"
]

for ref in references_list:
    p_ref = doc.add_paragraph()
    p_ref.paragraph_format.left_indent = Inches(0.3)
    p_ref.paragraph_format.first_line_indent = Inches(-0.3)
    p_ref.paragraph_format.space_after = Pt(3)
    r = p_ref.add_run(ref)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(9)

docx_path = Path("docs/Laporan_TK2_ANMEDSOS.docx")
doc.save(str(docx_path))
print(f"Saved docx to: {docx_path}")
